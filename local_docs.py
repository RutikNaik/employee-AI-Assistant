"""Local folder: load documents and rank them by keyword match."""
import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import config

SUPPORTED = {".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".md", ".csv"}


@dataclass
class Doc:
    path: Path
    name: str
    modified: str
    text: str
    truncated: bool = False


_cache: dict[tuple[str, float], str] = {}


def _pdf_text(path: Path) -> str:
    import pymupdf
    with pymupdf.open(path) as pdf:
        return "\n".join(page.get_text() for page in pdf)


def _docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(",".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        out.append(f"## Slide {i}")
        out += [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    return "\n".join(out)


READERS = {".pdf": _pdf_text, ".docx": _docx_text, ".xlsx": _xlsx_text, ".pptx": _pptx_text}


def _read(path: Path) -> str:
    reader = READERS.get(path.suffix.lower())
    if reader:
        return reader(path)
    return path.read_text(encoding="utf-8", errors="ignore")


def load_docs(folder: Path) -> tuple[list[Doc], list[str]]:
    """Returns (docs, errors). Unreadable files are reported, not skipped silently."""
    if not folder.is_dir():
        raise FileNotFoundError(f"Docs folder not found: {folder}")

    docs, errors = [], []
    for path in sorted(folder.rglob("*")):
        if path.suffix.lower() not in SUPPORTED or path.name.startswith(("~$", ".")):
            continue
        mtime = path.stat().st_mtime
        key = (str(path), mtime)
        if key not in _cache:
            try:
                _cache[key] = _read(path)
            except Exception as e:
                errors.append(f"{path.name}: {e}")
                continue
        text = _cache[key]
        docs.append(Doc(
            path=path,
            name=path.name,
            modified=datetime.fromtimestamp(mtime).isoformat(timespec="minutes"),
            text=text[: config.MAX_CHARS_PER_FILE],
            truncated=len(text) > config.MAX_CHARS_PER_FILE,
        ))
    return docs, errors


def fits_in_context(docs: list[Doc]) -> bool:
    return sum(len(d.text) for d in docs) <= config.MAX_TOTAL_CHARS


def rank(docs: list[Doc], keywords: list[str], top_k: int) -> list[Doc]:
    terms = [k.lower().strip() for k in keywords if k.strip()]

    def score(d: Doc) -> float:
        body, name = d.text.lower(), d.name.lower()
        return sum(math.log1p(body.count(t)) + 3 * (t in name) for t in terms)

    scored = sorted(((score(d), d) for d in docs), key=lambda x: x[0], reverse=True)
    return [d for s, d in scored if s > 0][:top_k]
